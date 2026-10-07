from datetime import datetime, timezone
from io import BytesIO
from uuid import uuid4

from openpyxl import load_workbook
from pypdf import PdfReader

from app.models.compliance_case import ComplianceCase
from app.models.task import Task
from app.models.verification_case import IdentityVerificationCase
from app.utils.enums import (
    ComplianceCasePriority,
    ComplianceCaseStatus,
    ComplianceCaseType,
    ReportExportFormat,
    ReportExportStatus,
    ReportType,
    SLAStatus,
    TaskPriority,
    TaskStatus,
    UserRole,
    VerificationStatus,
    VerificationType,
)
from tests.helpers import authenticate_client

REPORT_DATE = datetime(
    2035,
    1,
    15,
    12,
    0,
    tzinfo=timezone.utc,
)


def _auth_as_admin(
    client,
    create_test_user,
):
    admin = create_test_user(
        role=UserRole.ADMINISTRATOR,
        email=f"report-export-admin-{uuid4()}@example.com",
    )

    authenticate_client(
        client,
        admin,
    )

    return admin


def _auth_as_compliance_officer(
    client,
    create_test_user,
):
    officer = create_test_user(
        role=UserRole.COMPLIANCE_OFFICER,
        email=f"report-export-officer-{uuid4()}@example.com",
    )

    authenticate_client(
        client,
        officer,
    )

    return officer


def _auth_as_auditor(
    client,
    create_test_user,
):
    auditor = create_test_user(
        role=UserRole.AUDITOR,
        email=f"report-export-auditor-{uuid4()}@example.com",
    )

    authenticate_client(
        client,
        auditor,
    )

    return auditor


def test_customer_csv_report_is_generated_and_downloaded(
    client,
    create_test_user,
    create_test_customer,
    cleanup_test_customers,
    cleanup_report_exports,
):
    _auth_as_admin(
        client,
        create_test_user,
    )

    matching_customer = create_test_customer(
        first_name="Report",
        last_name="Customer",
        country_of_residence="REPORT-COUNTRY",
        email=f"report-customer-{uuid4()}@example.com",
        created_at=REPORT_DATE,
    )

    non_matching_customer = create_test_customer(
        first_name="Excluded",
        last_name="Customer",
        country_of_residence="OTHER-COUNTRY",
        email=f"excluded-customer-{uuid4()}@example.com",
        created_at=REPORT_DATE,
    )

    response = client.post(
        "/api/v1/reports/exports",
        json={
            "report_type": ReportType.CUSTOMER.value,
            "format": ReportExportFormat.CSV.value,
            "filters": {
                "start_date": "2035-01-15",
                "end_date": "2035-01-15",
                "country": "REPORT-COUNTRY",
            },
            "prefer_async": False,
        },
    )

    assert response.status_code == 201

    body = response.json()

    assert body["success"] is True

    export_data = body["data"]

    assert export_data["report_type"] == ReportType.CUSTOMER.value
    assert export_data["format"] == ReportExportFormat.CSV.value
    assert export_data["status"] == ReportExportStatus.COMPLETED.value
    assert export_data["row_count"] == 1
    assert export_data["filename"].endswith(".csv")
    assert export_data["content_type"] == "text/csv"
    assert export_data["file_size"] > 0
    assert export_data["download_url"] is not None

    export_id = export_data["id"]

    download_response = client.get(
        f"/api/v1/reports/exports/{export_id}/download",
    )

    assert download_response.status_code == 200
    assert download_response.headers["content-type"].startswith("text/csv")
    assert matching_customer.email.encode() in download_response.content
    assert non_matching_customer.email.encode() not in download_response.content


def test_customer_excel_report_is_generated_and_downloaded(
    client,
    create_test_user,
    create_test_customer,
    cleanup_test_customers,
    cleanup_report_exports,
):
    _auth_as_admin(
        client,
        create_test_user,
    )

    customer = create_test_customer(
        first_name="Excel",
        last_name="Customer",
        country_of_residence="EXCEL-COUNTRY",
        email=f"excel-customer-{uuid4()}@example.com",
        created_at=REPORT_DATE,
    )

    response = client.post(
        "/api/v1/reports/exports",
        json={
            "report_type": ReportType.CUSTOMER.value,
            "format": ReportExportFormat.EXCEL.value,
            "filters": {
                "start_date": "2035-01-15",
                "end_date": "2035-01-15",
                "country": "EXCEL-COUNTRY",
            },
            "prefer_async": False,
        },
    )

    assert response.status_code == 201

    export_data = response.json()["data"]

    assert export_data["status"] == ReportExportStatus.COMPLETED.value
    assert export_data["format"] == ReportExportFormat.EXCEL.value
    assert export_data["row_count"] == 1
    assert export_data["filename"].endswith(".xlsx")
    assert (
        export_data["content_type"]
        == "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
    )

    download_response = client.get(
        f"/api/v1/reports/exports/{export_data['id']}/download",
    )

    assert download_response.status_code == 200
    assert download_response.content.startswith(b"PK")

    workbook = load_workbook(
        BytesIO(download_response.content),
        read_only=True,
    )

    worksheet = workbook["Report"]

    rows = list(
        worksheet.iter_rows(
            values_only=True,
        )
    )

    assert rows[0][0] == "Customer ID"
    assert rows[0][1] == "Customer Name"
    assert rows[0][2] == "Email"

    assert str(customer.id) == str(rows[1][0])
    assert customer.email == rows[1][2]

    assert rows[1][9] is not None

    workbook.close()


def test_customer_pdf_report_is_generated_and_downloaded(
    client,
    create_test_user,
    create_test_customer,
    cleanup_test_customers,
    cleanup_report_exports,
):
    _auth_as_admin(
        client,
        create_test_user,
    )

    customer = create_test_customer(
        first_name="PDF",
        last_name="Customer",
        country_of_residence="PDF-COUNTRY",
        email=f"pdf-customer-{uuid4()}@example.com",
        created_at=REPORT_DATE,
    )

    response = client.post(
        "/api/v1/reports/exports",
        json={
            "report_type": ReportType.CUSTOMER.value,
            "format": ReportExportFormat.PDF.value,
            "filters": {
                "start_date": "2035-01-15",
                "end_date": "2035-01-15",
                "country": "PDF-COUNTRY",
            },
            "prefer_async": False,
        },
    )

    assert response.status_code == 201

    export_data = response.json()["data"]

    assert export_data["status"] == ReportExportStatus.COMPLETED.value
    assert export_data["format"] == ReportExportFormat.PDF.value
    assert export_data["row_count"] == 1
    assert export_data["filename"].endswith(".pdf")
    assert export_data["content_type"] == "application/pdf"

    download_response = client.get(
        f"/api/v1/reports/exports/{export_data['id']}/download",
    )

    assert download_response.status_code == 200
    assert download_response.headers["content-type"].startswith("application/pdf")
    assert download_response.content.startswith(b"%PDF")
    assert len(download_response.content) > 100

    reader = PdfReader(
        BytesIO(download_response.content),
    )

    pdf_text = "\n".join(page.extract_text() or "" for page in reader.pages)

    normalized_pdf_text = "".join(pdf_text.split())

    assert "PDFCustomer" in normalized_pdf_text
    assert "".join(str(customer.id).split()) in normalized_pdf_text
    assert "".join(customer.email.split()) in normalized_pdf_text


def test_verification_report_applies_filters(
    client,
    db_session,
    create_test_user,
    create_test_customer,
    cleanup_test_customers,
    cleanup_report_exports,
):
    _auth_as_admin(
        client,
        create_test_user,
    )

    matching_customer = create_test_customer(
        first_name="Verification",
        last_name="Match",
        country_of_residence="VERIFICATION-COUNTRY",
        email=f"verification-match-{uuid4()}@example.com",
        created_at=REPORT_DATE,
    )

    excluded_customer = create_test_customer(
        first_name="Verification",
        last_name="Excluded",
        country_of_residence="OTHER-COUNTRY",
        email=f"verification-excluded-{uuid4()}@example.com",
        created_at=REPORT_DATE,
    )

    db_session.add_all(
        [
            IdentityVerificationCase(
                customer_id=matching_customer.id,
                verification_type=VerificationType.IDENTITY,
                status=VerificationStatus.APPROVED,
                created_at=REPORT_DATE,
                completed_at=REPORT_DATE,
            ),
            IdentityVerificationCase(
                customer_id=excluded_customer.id,
                verification_type=VerificationType.IDENTITY,
                status=VerificationStatus.REJECTED,
                created_at=REPORT_DATE,
                completed_at=REPORT_DATE,
            ),
        ]
    )

    db_session.commit()

    response = client.post(
        "/api/v1/reports/exports",
        json={
            "report_type": ReportType.VERIFICATION.value,
            "format": ReportExportFormat.CSV.value,
            "filters": {
                "start_date": "2035-01-15",
                "end_date": "2035-01-15",
                "country": "VERIFICATION-COUNTRY",
                "verification_status": VerificationStatus.APPROVED.value,
                "verification_type": VerificationType.IDENTITY.value,
            },
            "prefer_async": False,
        },
    )

    assert response.status_code == 201

    export_data = response.json()["data"]

    assert export_data["status"] == ReportExportStatus.COMPLETED.value
    assert export_data["row_count"] == 1

    download_response = client.get(
        f"/api/v1/reports/exports/{export_data['id']}/download",
    )

    assert download_response.status_code == 200

    content = download_response.content.decode("utf-8-sig")

    assert str(matching_customer.id) in content
    assert str(excluded_customer.id) not in content
    assert "Verification Match" in content
    assert "Verification Excluded" not in content
    assert "APPROVED" in content
    assert "REJECTED" not in content


def test_compliance_report_applies_status_and_country_filters(
    client,
    db_session,
    create_test_user,
    create_test_customer,
    cleanup_test_customers,
    cleanup_compliance_cases,
    cleanup_report_exports,
):
    _auth_as_admin(
        client,
        create_test_user,
    )

    matching_customer = create_test_customer(
        first_name="Compliance",
        last_name="Match",
        country_of_residence="COMPLIANCE-COUNTRY",
        email=f"compliance-match-{uuid4()}@example.com",
        created_at=REPORT_DATE,
    )

    excluded_customer = create_test_customer(
        first_name="Compliance",
        last_name="Excluded",
        country_of_residence="OTHER-COUNTRY",
        email=f"compliance-excluded-{uuid4()}@example.com",
        created_at=REPORT_DATE,
    )

    db_session.add_all(
        [
            ComplianceCase(
                customer_id=matching_customer.id,
                case_type=ComplianceCaseType.CUSTOMER_REVIEW,
                priority=ComplianceCasePriority.HIGH,
                status=ComplianceCaseStatus.OPEN,
                created_at=REPORT_DATE,
            ),
            ComplianceCase(
                customer_id=excluded_customer.id,
                case_type=ComplianceCaseType.CUSTOMER_REVIEW,
                priority=ComplianceCasePriority.HIGH,
                status=ComplianceCaseStatus.CLOSED,
                created_at=REPORT_DATE,
                closed_at=REPORT_DATE,
                resolution_reason="Excluded test case",
            ),
        ]
    )

    db_session.commit()

    response = client.post(
        "/api/v1/reports/exports",
        json={
            "report_type": ReportType.COMPLIANCE.value,
            "format": ReportExportFormat.CSV.value,
            "filters": {
                "start_date": "2035-01-15",
                "end_date": "2035-01-15",
                "country": "COMPLIANCE-COUNTRY",
                "compliance_status": ComplianceCaseStatus.OPEN.value,
            },
            "prefer_async": False,
        },
    )

    assert response.status_code == 201

    export_data = response.json()["data"]

    assert export_data["status"] == ReportExportStatus.COMPLETED.value
    assert export_data["row_count"] == 1

    download_response = client.get(
        f"/api/v1/reports/exports/{export_data['id']}/download",
    )

    assert download_response.status_code == 200

    content = download_response.content.decode("utf-8-sig")

    assert str(matching_customer.id) in content
    assert str(excluded_customer.id) not in content
    assert "Compliance Match" in content
    assert "Compliance Excluded" not in content
    assert "OPEN" in content
    assert "CLOSED" not in content


def test_operational_performance_report_contains_tasks(
    client,
    db_session,
    create_test_user,
    cleanup_tasks,
    cleanup_report_exports,
):
    admin = _auth_as_admin(
        client,
        create_test_user,
    )

    task = Task(
        title="Report export test task",
        description="Task created for reporting export test.",
        assigned_to=admin.id,
        priority=TaskPriority.HIGH,
        status=TaskStatus.COMPLETED,
        due_date=REPORT_DATE,
        completed_at=REPORT_DATE,
        sla_status=SLAStatus.COMPLETED,
        created_at=REPORT_DATE,
    )

    db_session.add(task)
    db_session.commit()

    response = client.post(
        "/api/v1/reports/exports",
        json={
            "report_type": ReportType.OPERATIONAL_PERFORMANCE.value,
            "format": ReportExportFormat.CSV.value,
            "filters": {
                "start_date": "2035-01-15",
                "end_date": "2035-01-15",
                "task_status": TaskStatus.COMPLETED.value,
            },
            "prefer_async": False,
        },
    )

    assert response.status_code == 201

    export_data = response.json()["data"]

    assert export_data["status"] == ReportExportStatus.COMPLETED.value
    assert export_data["row_count"] == 1

    download_response = client.get(
        f"/api/v1/reports/exports/{export_data['id']}/download",
    )

    assert download_response.status_code == 200

    content = download_response.content.decode("utf-8-sig")

    assert "TASK" in content
    assert str(task.id) in content
    assert "Report export test task" in content
    assert "COMPLETED" in content


def test_large_report_can_be_queued_for_background_processing(
    client,
    create_test_user,
    create_test_customer,
    cleanup_test_customers,
    cleanup_report_exports,
):
    _auth_as_admin(
        client,
        create_test_user,
    )

    create_test_customer(
        first_name="Async",
        last_name="Customer",
        country_of_residence="ASYNC-COUNTRY",
        email=f"async-customer-{uuid4()}@example.com",
        created_at=REPORT_DATE,
    )

    response = client.post(
        "/api/v1/reports/exports",
        json={
            "report_type": ReportType.CUSTOMER.value,
            "format": ReportExportFormat.CSV.value,
            "filters": {
                "start_date": "2035-01-15",
                "end_date": "2035-01-15",
                "country": "ASYNC-COUNTRY",
            },
            "prefer_async": True,
        },
    )

    assert response.status_code == 202

    export_data = response.json()["data"]

    assert export_data["status"] == ReportExportStatus.REQUESTED.value
    assert export_data["row_count"] == 1
    assert export_data["download_url"] is None

    status_response = client.get(
        f"/api/v1/reports/exports/{export_data['id']}",
    )

    assert status_response.status_code == 200

    status_data = status_response.json()["data"]

    assert status_data["id"] == export_data["id"]
    assert status_data["status"] == ReportExportStatus.REQUESTED.value


def test_report_export_history_is_stored(
    client,
    create_test_user,
    create_test_customer,
    cleanup_test_customers,
    cleanup_report_exports,
):
    _auth_as_admin(
        client,
        create_test_user,
    )

    create_test_customer(
        first_name="History",
        last_name="Customer",
        country_of_residence="HISTORY-COUNTRY",
        email=f"history-customer-{uuid4()}@example.com",
        created_at=REPORT_DATE,
    )

    first_response = client.post(
        "/api/v1/reports/exports",
        json={
            "report_type": ReportType.CUSTOMER.value,
            "format": ReportExportFormat.CSV.value,
            "filters": {
                "start_date": "2035-01-15",
                "end_date": "2035-01-15",
                "country": "HISTORY-COUNTRY",
            },
            "prefer_async": False,
        },
    )

    assert first_response.status_code == 201

    first_export_id = first_response.json()["data"]["id"]

    second_response = client.post(
        "/api/v1/reports/exports",
        json={
            "report_type": ReportType.CUSTOMER.value,
            "format": ReportExportFormat.EXCEL.value,
            "filters": {
                "start_date": "2035-01-15",
                "end_date": "2035-01-15",
                "country": "HISTORY-COUNTRY",
            },
            "prefer_async": False,
        },
    )

    assert second_response.status_code == 201

    second_export_id = second_response.json()["data"]["id"]

    history_response = client.get(
        "/api/v1/reports/exports",
        params={
            "page": 1,
            "page_size": 100,
        },
    )

    assert history_response.status_code == 200

    history_data = history_response.json()["data"]

    assert history_data["total"] >= 2

    export_ids = {item["id"] for item in history_data["exports"]}

    assert first_export_id in export_ids
    assert second_export_id in export_ids


def test_auditor_can_view_report_history(
    client,
    create_test_user,
    create_test_customer,
    cleanup_test_customers,
    cleanup_report_exports,
):
    admin = _auth_as_admin(
        client,
        create_test_user,
    )

    create_test_customer(
        country_of_residence="AUDITOR-COUNTRY",
        email=f"auditor-report-{uuid4()}@example.com",
        created_at=REPORT_DATE,
    )

    create_response = client.post(
        "/api/v1/reports/exports",
        json={
            "report_type": ReportType.CUSTOMER.value,
            "format": ReportExportFormat.CSV.value,
            "filters": {
                "start_date": "2035-01-15",
                "end_date": "2035-01-15",
                "country": "AUDITOR-COUNTRY",
            },
            "prefer_async": False,
        },
    )

    assert create_response.status_code == 201

    export_id = create_response.json()["data"]["id"]

    assert admin is not None

    _auth_as_auditor(
        client,
        create_test_user,
    )

    history_response = client.get(
        "/api/v1/reports/exports",
    )

    assert history_response.status_code == 200

    export_ids = {item["id"] for item in history_response.json()["data"]["exports"]}

    assert export_id in export_ids


def test_auditor_can_download_completed_report(
    client,
    create_test_user,
    create_test_customer,
    cleanup_test_customers,
    cleanup_report_exports,
):
    _auth_as_admin(
        client,
        create_test_user,
    )

    create_test_customer(
        first_name="Auditor",
        last_name="Download",
        country_of_residence="AUDITOR-DOWNLOAD-COUNTRY",
        email=f"auditor-download-{uuid4()}@example.com",
        created_at=REPORT_DATE,
    )

    create_response = client.post(
        "/api/v1/reports/exports",
        json={
            "report_type": ReportType.CUSTOMER.value,
            "format": ReportExportFormat.CSV.value,
            "filters": {
                "start_date": "2035-01-15",
                "end_date": "2035-01-15",
                "country": "AUDITOR-DOWNLOAD-COUNTRY",
            },
            "prefer_async": False,
        },
    )

    assert create_response.status_code == 201

    export_id = create_response.json()["data"]["id"]

    _auth_as_auditor(
        client,
        create_test_user,
    )

    download_response = client.get(
        f"/api/v1/reports/exports/{export_id}/download",
    )

    assert download_response.status_code == 200
    assert download_response.content.decode("utf-8-sig").startswith("Customer ID")


def test_compliance_officer_can_generate_report(
    client,
    create_test_user,
    create_test_customer,
    cleanup_test_customers,
    cleanup_report_exports,
):
    _auth_as_compliance_officer(
        client,
        create_test_user,
    )

    create_test_customer(
        first_name="Compliance",
        last_name="Officer",
        country_of_residence="OFFICER-COUNTRY",
        email=f"officer-report-{uuid4()}@example.com",
        created_at=REPORT_DATE,
    )

    response = client.post(
        "/api/v1/reports/exports",
        json={
            "report_type": ReportType.CUSTOMER.value,
            "format": ReportExportFormat.CSV.value,
            "filters": {
                "start_date": "2035-01-15",
                "end_date": "2035-01-15",
                "country": "OFFICER-COUNTRY",
            },
            "prefer_async": False,
        },
    )

    assert response.status_code == 201

    data = response.json()["data"]

    assert data["status"] == ReportExportStatus.COMPLETED.value


def test_reviewer_cannot_generate_report_export(
    client,
    create_test_user,
    cleanup_report_exports,
):
    reviewer = create_test_user(
        role=UserRole.REVIEWER,
        email=f"report-export-reviewer-{uuid4()}@example.com",
    )

    authenticate_client(
        client,
        reviewer,
    )

    response = client.post(
        "/api/v1/reports/exports",
        json={
            "report_type": ReportType.CUSTOMER.value,
            "format": ReportExportFormat.CSV.value,
            "filters": {},
            "prefer_async": False,
        },
    )

    assert response.status_code == 403


def test_report_export_rejects_invalid_date_range(
    client,
    create_test_user,
    cleanup_report_exports,
):
    _auth_as_admin(
        client,
        create_test_user,
    )

    response = client.post(
        "/api/v1/reports/exports",
        json={
            "report_type": ReportType.CUSTOMER.value,
            "format": ReportExportFormat.CSV.value,
            "filters": {
                "start_date": "2035-01-20",
                "end_date": "2035-01-01",
            },
            "prefer_async": False,
        },
    )

    assert response.status_code == 422


def test_report_export_returns_not_found_for_unknown_export(
    client,
    create_test_user,
):
    _auth_as_admin(
        client,
        create_test_user,
    )

    response = client.get(
        f"/api/v1/reports/exports/{uuid4()}",
    )

    assert response.status_code == 404
